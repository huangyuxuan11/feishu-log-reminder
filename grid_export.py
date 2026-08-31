#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
网格长拜访日志 - 每周一自动导出上周记录为 Excel 并飞书推送
（设计运行于 GitHub Actions 云端，复用飞书 OpenAPI，无需 lark-cli）

数据源: 飞书多维表格「网格长拜访日志」
  base_token = F3b7b4E10aDtznsYf3IcLUtVnKe
  table_id   = tblA2vopNSsC1WSC

导出口径: 周报用得到的 12 字段，随访日期置首列，单 sheet；
          对"沟通内容"做"隐去纯时间段残留"轻量清洗。

触发: 由飞书多维表格自动化（每周一 08:00 北京时间）经 workflow_dispatch 触发。
环境变量(放 GitHub Secrets): FEISHU_APP_ID / FEISHU_APP_SECRET / FEISHU_USER_OPEN_ID
"""
import os
import io
import re
import json
import datetime
import requests

GRID_APP = "F3b7b4E10aDtznsYf3IcLUtVnKe"
GRID_TABLE = "tblA2vopNSsC1WSC"
FEISHU_BASE = "https://open.feishu.cn/open-apis"
BJT = datetime.timezone(datetime.timedelta(hours=8))

# 导出字段（随访日期置首列），与周报读取口径一致
EXPORT_FIELDS = [
    "随访日期", "所属网格/小组", "随访人", "其他随访人", "时间",
    "客户经理", "集团名称", "是否要客集团", "拜访对象",
    "系统中关键人", "系统中是否真实关键人", "沟通内容",
]

TIME_RANGE_RE = re.compile(
    r"\d{1,2}[:：]\d{2}\s*[~\u301c\uff5e\-]\s*\d{1,2}[:：]\d{2}"
    r"|\b\d{4}\s*[\-\u2013\u2014~\u301c\uff5e]\s*\d{4}\b"
)
_DATE_RE = re.compile(r"(\d{4}-\d{2}-\d{2})")


def clean_text(s):
    if not s:
        return s
    s = TIME_RANGE_RE.sub("", s)
    s = re.sub(r"[、，,]\s*[、，,]+", "，", s)
    s = re.sub(r"\s{2,}", " ", s)
    return s.strip(" 、，,.\u3002").strip()


def get_tenant_token(app_id, app_secret):
    url = f"{FEISHU_BASE}/auth/v3/tenant_access_token/internal"
    r = requests.post(url, json={"app_id": app_id, "app_secret": app_secret}, timeout=15)
    data = r.json()
    if data.get("code") != 0:
        raise RuntimeError(f"获取 tenant_access_token 失败: {data}")
    return data["tenant_access_token"]


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, list):
        return "，".join(str(x) for x in v)
    if isinstance(v, dict):
        for k in ("date", "start"):
            if k in v:
                return str(v[k])
        return str(v)
    return str(v)


def _date_only(v):
    """从飞书日期字段值中提取 YYYY-MM-DD（兼容 字符串/时间戳/dict/list）。"""
    if isinstance(v, dict):
        for k in ("date", "start"):
            if k in v:
                v = v[k]
                break
    if isinstance(v, list):
        v = v[0] if v else ""
    if isinstance(v, bool):
        pass
    elif isinstance(v, (int, float)):
        try:
            return datetime.datetime.fromtimestamp(v / 1000, tz=BJT).strftime("%Y-%m-%d")
        except (ValueError, OSError, OverflowError):
            return str(v)
    s = v if isinstance(v, str) else _norm(v)
    if isinstance(s, str):
        s = s.strip()
        m = _DATE_RE.search(s)
        return m.group(1) if m else s
    return str(v)


def list_all_records(app_token, table_id, token, page_size=100):
    url = f"{FEISHU_BASE}/bitable/v1/apps/{app_token}/tables/{table_id}/records"
    out, page_token = [], ""
    while True:
        params = {"page_size": page_size}
        if page_token:
            params["page_token"] = page_token
        r = requests.get(url, headers={"Authorization": f"Bearer {token}"}, params=params, timeout=20)
        data = r.json()
        if data.get("code") != 0:
            raise RuntimeError(f"读取记录失败: {data}")
        out.extend(data.get("data", {}).get("items", []))
        if not data.get("data", {}).get("has_more") or not data.get("data", {}).get("page_token"):
            break
        page_token = data["data"]["page_token"]
    return out


def last_week_range(now=None):
    now = now or datetime.datetime.now(BJT)
    today = now.date()
    last_mon = today - datetime.timedelta(days=today.weekday() + 7)
    last_sun = last_mon + datetime.timedelta(days=6)
    return last_mon, last_sun


def build_excel(records, start, end):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表单"
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(EXPORT_FIELDS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for r, it in enumerate(records, 2):
        f = it.get("fields", {})
        row_vals = []
        for name in EXPORT_FIELDS:
            v = f.get(name, "")
            if name == "沟通内容":
                v = clean_text(_norm(v))
            elif name == "随访日期":
                v = _date_only(v)
            else:
                v = _norm(v)
            row_vals.append(v)
        for c, v in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    widths = [13, 16, 10, 12, 14, 12, 22, 12, 14, 14, 18, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def upload_file(token, file_bytes, file_name):
    url = f"{FEISHU_BASE}/im/v1/files"
    files = {"file": (file_name, file_bytes,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"file_name": file_name, "file_type": "xlsx"}
    r = requests.post(url, headers={"Authorization": f"Bearer {token}"}, files=files, data=data, timeout=30)
    resp = r.json()
    if resp.get("code") != 0:
        raise RuntimeError(f"文件上传失败: {resp}")
    return resp["data"]["file_key"]


def send_file_message(token, open_id, file_key, receive_id_type="open_id"):
    url = f"{FEISHU_BASE}/im/v1/messages?receive_id_type={receive_id_type}"
    body = {"receive_id": open_id, "msg_type": "file", "content": json.dumps({"file_key": file_key})}
    r = requests.post(url, headers={"Authorization": f"Bearer {token}"}, json=body, timeout=20)
    return r.json()


def send_text_message(token, open_id, text, receive_id_type="open_id"):
    url = f"{FEISHU_BASE}/im/v1/messages?receive_id_type={receive_id_type}"
    body = {"receive_id": open_id, "msg_type": "text", "content": json.dumps({"text": text})}
    r = requests.post(url, headers={"Authorization": f"Bearer {token}"}, json=body, timeout=20)
    return r.json()


def main():
    app_id = os.environ.get("FEISHU_APP_ID")
    app_secret = os.environ.get("FEISHU_APP_SECRET")
    open_id = os.environ.get("FEISHU_USER_OPEN_ID")
    if not app_id or not app_secret or not open_id:
        raise SystemExit("缺少环境变量 FEISHU_APP_ID / FEISHU_APP_SECRET / FEISHU_USER_OPEN_ID")

    start, end = last_week_range()
    print(f"[info] 导出区间: {start} ~ {end}")

    token = get_tenant_token(app_id, app_secret)
    items = list_all_records(GRID_APP, GRID_TABLE, token)
    print(f"[info] 全量记录数: {len(items)}")

    matched = []
    for it in items:
        d = _date_only(it.get("fields", {}).get("随访日期"))
        try:
            rd = datetime.datetime.strptime(d, "%Y-%m-%d").date()
        except Exception:
            continue
        if start <= rd <= end:
            matched.append(it)
    matched.sort(key=lambda it: (it.get("fields", {}).get("随访日期", ""),
                                  it.get("fields", {}).get("所属网格/小组", "")))
    print(f"[info] 命中记录数: {len(matched)}")

    file_name = f"网格长拜访日志{start.strftime('%Y%m%d')}-{end.strftime('%Y%m%d')}.xlsx"
    buf = build_excel(matched, start, end)
    payload = buf.getvalue()
    # 本地留一份（CI 中即仓库工作区文件，便于排障）
    with open(file_name, "wb") as fh:
        fh.write(payload)
    print(f"[info] 本地文件: {file_name}")

    file_key = upload_file(token, payload, file_name)
    print(f"[info] 文件已上传 file_key={file_key}")

    intro = (f"📊 网格长拜访日志（{start.strftime('%Y%m%d')}-{end.strftime('%Y%m%d')}）"
             f"共 {len(matched)} 条，见附件Excel。")
    send_text_message(token, open_id, intro)
    resp = send_file_message(token, open_id, file_key)
    print("file 返回:", resp)
    print(f"✅ 已发送 {file_name}")


if __name__ == "__main__":
    main()
